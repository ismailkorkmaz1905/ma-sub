"""Create the small, uncertainty-only Turkish review workbook."""

from __future__ import annotations

from collections import defaultdict
from collections.abc import Mapping, Sequence
from pathlib import Path
import math
import os
import tempfile
from typing import Any

try:  # Package and Colab flat-module imports are both supported.
    from .srt import format_timestamp
except ImportError:  # pragma: no cover - exercised in Colab notebook mode
    from srt import format_timestamp


REVIEW_COLUMNS: tuple[str, ...] = (
    "Öncelik",
    "SRT Blok",
    "Zaman",
    "Final Türkçe",
    "Whisper",
    "YouTube Auto",
    "Final Endonezce",
    "Kontrol Notu",
    "Kullanıcı Düzeltmesi",
    "Durum",
)


class ReviewWorkbookError(RuntimeError):
    """Raised when a review workbook cannot be produced or verified."""


def _records_by_uid(value: Any) -> Mapping[str, Mapping[str, Any]]:
    if hasattr(value, "records_by_uid"):
        value = getattr(value, "records_by_uid")
    elif hasattr(value, "validated_by_uid"):
        value = getattr(value, "validated_by_uid")
    if isinstance(value, Mapping):
        output: dict[str, Mapping[str, Any]] = {}
        for key, record in value.items():
            if isinstance(record, Mapping):
                output[str(key)] = record
        return output
    if isinstance(value, Sequence) and not isinstance(value, (str, bytes)):
        output = {}
        for record in value:
            if not isinstance(record, Mapping):
                continue
            uid = record.get("block_uid")
            if isinstance(uid, str) and uid and uid not in output:
                output[uid] = record
        return output
    raise TypeError("translations_by_uid must be a mapping or sequence")


def _issue_list(value: Any) -> list[Mapping[str, Any]]:
    if isinstance(value, Mapping) and "issues" in value:
        value = value.get("issues")
    if value is None:
        return []
    if isinstance(value, Sequence) and not isinstance(value, (str, bytes)):
        return [item for item in value if isinstance(item, Mapping)]
    raise TypeError("issues must be a QA report or a sequence of issue objects")


def _risk_flags(block: Mapping[str, Any]) -> list[str]:
    value = block.get("risk_flags", ())
    if isinstance(value, str):
        return [value] if value.strip() else []
    if isinstance(value, Sequence) and not isinstance(value, (str, bytes)):
        return [str(flag) for flag in value if str(flag).strip()]
    return []


def _priority(
    record: Mapping[str, Any],
    flags: Sequence[str],
    block_issues: Sequence[Mapping[str, Any]],
) -> tuple[int, str]:
    error_count = sum(issue.get("severity", "error") != "warning" for issue in block_issues)
    warning_count = len(block_issues) - error_count
    if error_count:
        return 100 + error_count, "Yüksek"
    if record.get("review_required") is True:
        return 90, "Yüksek"
    if flags:
        return 60 + min(len(flags), 10), "Orta"
    if warning_count:
        return 40 + min(warning_count, 10), "Orta"
    return 20, "Düşük"


def _note_text(
    record: Mapping[str, Any],
    flags: Sequence[str],
    block_issues: Sequence[Mapping[str, Any]],
) -> str:
    notes: list[str] = []
    translator_note = record.get("note")
    if isinstance(translator_note, str) and translator_note.strip():
        notes.append(translator_note.strip())
    if flags:
        notes.append("Risk: " + ", ".join(flags))
    for issue in block_issues:
        message = issue.get("message") or issue.get("code")
        if message:
            notes.append(str(message).strip())
    # Preserve order while removing repeated diagnostics.
    return " | ".join(dict.fromkeys(note for note in notes if note))


def collect_review_rows(
    blocks: Sequence[Mapping[str, Any]],
    translations_by_uid: Any,
    issues: Any,
    *,
    max_fraction: float | None = 0.05,
) -> list[dict[str, Any]]:
    """Select and format only genuinely uncertain rows.

    Rows are ranked deterministically. ``max_fraction`` caps the total review
    sheet at the highest-priority fraction of episode blocks, except that every
    explicit translator review request is always retained. Use ``None`` to
    retain all flagged rows. A clean episode intentionally produces a
    header-only sheet.
    """

    if max_fraction is not None and not (0 < max_fraction <= 1):
        raise ValueError("max_fraction must be in the interval (0, 1] or None")
    records = _records_by_uid(translations_by_uid)
    grouped_issues: dict[str, list[Mapping[str, Any]]] = defaultdict(list)
    for issue in _issue_list(issues):
        uid = issue.get("block_uid")
        if isinstance(uid, str) and uid:
            grouped_issues[uid].append(issue)

    candidates: list[tuple[bool, int, int, dict[str, Any]]] = []
    for position, block in enumerate(blocks, start=1):
        uid = block.get("block_uid")
        if not isinstance(uid, str) or uid not in records:
            continue
        record = records[uid]
        flags = _risk_flags(block)
        block_issues = grouped_issues.get(uid, [])
        translator_note = record.get("note")
        uncertain = (
            record.get("review_required") is True
            or bool(flags)
            or bool(block_issues)
            or (isinstance(translator_note, str) and bool(translator_note.strip()))
        )
        if not uncertain:
            continue
        score, label = _priority(record, flags, block_issues)
        index = block.get("block_index", position)
        try:
            index = int(index)
        except (TypeError, ValueError):
            index = position
        start_ms = block.get("start_ms")
        end_ms = block.get("end_ms")
        if isinstance(start_ms, int) and isinstance(end_ms, int) and start_ms >= 0 and end_ms >= 0:
            timing = f"{format_timestamp(start_ms)} --> {format_timestamp(end_ms)}"
        else:
            timing = f"{start_ms!r} --> {end_ms!r}"
        row = {
            "Öncelik": label,
            "SRT Blok": index,
            "Zaman": timing,
            "Final Türkçe": str(record.get("tr_final") or ""),
            "Whisper": str(block.get("primary_text") or ""),
            "YouTube Auto": str(block.get("youtube_text") or ""),
            "Final Endonezce": str(record.get("id_final") or ""),
            "Kontrol Notu": _note_text(record, flags, block_issues),
            "Kullanıcı Düzeltmesi": "",
            "Durum": "Bekliyor",
        }
        # An explicit translator review request is never discarded by the
        # heuristic sheet-size target.
        mandatory = record.get("review_required") is True
        candidates.append((mandatory, score, index, row))

    candidates.sort(key=lambda item: (-item[1], item[2]))
    if max_fraction is not None and candidates:
        total_capacity = max(1, math.ceil(len(blocks) * max_fraction))
        mandatory_rows = [item for item in candidates if item[0]]
        heuristic_capacity = max(0, total_capacity - len(mandatory_rows))
        heuristic_rows = [item for item in candidates if not item[0]][
            :heuristic_capacity
        ]
        candidates = mandatory_rows + heuristic_rows
    # The worksheet is easiest to use in chronological/SRT order.
    candidates.sort(key=lambda item: item[2])
    return [row for _, _, _, row in candidates]


def _save_atomic_and_verify(workbook: Any, destination: Path, expected_rows: int) -> None:
    from openpyxl import load_workbook

    destination.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{destination.stem}.", suffix=".tmp.xlsx", dir=str(destination.parent)
    )
    os.close(descriptor)
    try:
        workbook.save(temporary_name)
        # Re-open the exact bytes which will be published. This catches partial
        # ZIP/XLSX writes and verifies the Turkish header contract.
        check = load_workbook(temporary_name, read_only=True, data_only=False)
        try:
            if check.sheetnames != ["Kontrol"]:
                raise ReviewWorkbookError(
                    f"Review workbook sheets differ: {check.sheetnames!r}"
                )
            sheet = check["Kontrol"]
            headers = tuple(sheet.cell(1, column).value for column in range(1, 11))
            if headers != REVIEW_COLUMNS:
                raise ReviewWorkbookError(
                    f"Review workbook headers differ: {headers!r}"
                )
            if sheet.max_row != expected_rows + 1:
                raise ReviewWorkbookError(
                    "Review workbook row count differs: "
                    f"expected {expected_rows + 1}, got {sheet.max_row}"
                )
        finally:
            check.close()
        with open(temporary_name, "rb") as handle:
            os.fsync(handle.fileno())
        os.replace(temporary_name, destination)
    except BaseException:
        try:
            os.unlink(temporary_name)
        except FileNotFoundError:
            pass
        raise


def create_review_xlsx(
    path: str | os.PathLike[str],
    blocks: Sequence[Mapping[str, Any]],
    translations_by_uid: Any,
    issues: Any,
    max_fraction: float | None = 0.05,
) -> Path:
    """Create an atomic, styled, Colab-compatible ``Kontrol`` workbook."""

    try:
        from openpyxl import Workbook
        from openpyxl.formatting.rule import FormulaRule
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError as exc:  # pragma: no cover - dependency is installed in Colab
        raise ReviewWorkbookError(
            "openpyxl is required to create the review workbook"
        ) from exc

    destination = Path(path)
    if destination.suffix.lower() != ".xlsx":
        raise ReviewWorkbookError("Review workbook path must end with .xlsx")
    rows = collect_review_rows(
        blocks,
        translations_by_uid,
        issues,
        max_fraction=max_fraction,
    )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Kontrol"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.append(REVIEW_COLUMNS)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin_gray = Side(style="thin", color="D9E2F3")
    body_border = Border(bottom=thin_gray)
    alternate_fill = PatternFill("solid", fgColor="F5F9FD")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 30

    for row_number, row in enumerate(rows, start=2):
        for column_number, column_name in enumerate(REVIEW_COLUMNS, start=1):
            cell = sheet.cell(row_number, column_number)
            value = row[column_name]
            cell.value = value
            # Prevent subtitle text beginning with =, +, -, or @ from becoming
            # an executable spreadsheet formula.
            if isinstance(value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = body_border
            if row_number % 2 == 0:
                cell.fill = alternate_fill
        sheet.cell(row_number, 1).alignment = Alignment(
            horizontal="center", vertical="top"
        )
        sheet.cell(row_number, 2).alignment = Alignment(
            horizontal="center", vertical="top"
        )
        sheet.row_dimensions[row_number].height = 48

    last_row = max(1, len(rows) + 1)
    sheet.auto_filter.ref = f"A1:J{last_row}"
    sheet.print_title_rows = "1:1"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25

    widths = (11, 11, 27, 42, 42, 42, 42, 60, 42, 15)
    for column_number, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column_number)].width = width

    if rows:
        validation = DataValidation(
            type="list",
            formula1='"Bekliyor,Kontrol Edildi,Düzeltildi"',
            allow_blank=False,
        )
        validation.error = "Listeden geçerli bir durum seçin."
        validation.errorTitle = "Geçersiz durum"
        validation.prompt = "Kontrol durumunu seçin."
        validation.promptTitle = "Durum"
        sheet.add_data_validation(validation)
        validation.add(f"J2:J{last_row}")

        high_fill = PatternFill("solid", fgColor="F4CCCC")
        medium_fill = PatternFill("solid", fgColor="FCE5CD")
        low_fill = PatternFill("solid", fgColor="FFF2CC")
        sheet.conditional_formatting.add(
            f"A2:A{last_row}", FormulaRule(formula=['A2="Yüksek"'], fill=high_fill)
        )
        sheet.conditional_formatting.add(
            f"A2:A{last_row}", FormulaRule(formula=['A2="Orta"'], fill=medium_fill)
        )
        sheet.conditional_formatting.add(
            f"A2:A{last_row}", FormulaRule(formula=['A2="Düşük"'], fill=low_fill)
        )

    _save_atomic_and_verify(workbook, destination, len(rows))
    workbook.close()
    return destination


__all__ = [
    "REVIEW_COLUMNS",
    "ReviewWorkbookError",
    "collect_review_rows",
    "create_review_xlsx",
]

