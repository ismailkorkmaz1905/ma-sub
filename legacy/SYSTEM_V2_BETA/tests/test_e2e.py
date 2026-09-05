from __future__ import annotations

import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from helpers import synthetic_transcription, translation_records, write_translated_zip
from src.batches import create_translation_pack, validate_translation_pack
from src.review import REVIEW_COLUMNS, create_review_xlsx
from src.schema import build_episode_schema
from src.segment import segment_transcription
from src.srt import build_entries, parse_srt, write_srt
from src.subtitle_qa import assert_final_qa, run_subtitle_qa
from src.translation_validation import load_and_validate_translated_zip


class SyntheticWorkflowTests(unittest.TestCase):
    def test_prepare_pack_to_translated_zip_to_final_srt_and_review(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            transcription = synthetic_transcription(10)
            segmentation = segment_transcription(
                transcription, episode=12, prepare_dir=root / "prepare"
            )
            resumed = segment_transcription(
                transcription, episode=12, prepare_dir=root / "prepare"
            )
            self.assertTrue(segmentation.report.valid)
            self.assertFalse(segmentation.resumed)
            self.assertTrue(resumed.resumed)
            self.assertEqual(segmentation.blocks, resumed.blocks)
            self.assertEqual(len(segmentation.blocks), 10)

            schema = build_episode_schema(segmentation.blocks, episode=12)
            records = translation_records(schema)
            records[3]["review_required"] = True
            records[3]["note"] = "Bağlama göre hitap biçimini kontrol et."
            pack = root / "Muhtemel Ask 12.Bolum_TRANSLATION_PACK.zip"
            translated = root / "Muhtemel Ask 12.Bolum_TRANSLATED.zip"
            tr_path = root / "Muhtemel Ask 12.Bolum.tr-final.srt"
            id_path = root / "Muhtemel Ask 12.Bolum.id-final.srt"
            review_path = root / "Muhtemel Ask 12.Bolum_review.xlsx"

            manifest = create_translation_pack(schema, pack)
            self.assertEqual(validate_translation_pack(pack, expected_schema=schema), manifest)
            with zipfile.ZipFile(pack, "r") as archive:
                self.assertFalse(
                    any(
                        Path(name).suffix.lower() in {".mkv", ".mp4", ".flac", ".wav"}
                        for name in archive.namelist()
                    )
                )
                glossary = json.loads(archive.read("glossary.json"))
                self.assertEqual(
                    set(glossary),
                    {
                        "canonical_names",
                        "forbidden_name_variants",
                        "source_name_variants",
                        "religious_terms",
                        "allah_only_semoga_is_invalid",
                    },
                )
                self.assertIn(
                    "Selam",
                    glossary["source_name_variants"]["Selim"],
                )

            write_translated_zip(translated, schema, records)
            validated = load_and_validate_translated_zip(
                schema,
                translated,
                input_manifest=manifest,
            )
            self.assertTrue(validated.ok)
            self.assertEqual(validated.report.positional_translation_mismatch_count, 0)

            qa_report = run_subtitle_qa(schema["blocks"], validated)
            assert_final_qa(qa_report)
            tr_entries = build_entries(schema["blocks"], validated.records_by_uid, "tr")
            id_entries = build_entries(schema["blocks"], validated.records_by_uid, "id")
            write_srt(tr_path, tr_entries)
            write_srt(id_path, id_entries)
            create_review_xlsx(
                review_path,
                schema["blocks"],
                validated.records_by_uid,
                qa_report,
                max_fraction=0.20,
            )

            parsed_tr = parse_srt(tr_path)
            parsed_id = parse_srt(id_path)
            self.assertEqual(len(parsed_tr), schema["block_count"])
            self.assertEqual(len(parsed_id), schema["block_count"])
            self.assertEqual(
                [(entry.start_ms, entry.end_ms) for entry in parsed_tr],
                [(entry.start_ms, entry.end_ms) for entry in parsed_id],
            )
            workbook = load_workbook(review_path, read_only=True)
            try:
                self.assertEqual(workbook.sheetnames, ["Kontrol"])
                sheet = workbook["Kontrol"]
                self.assertEqual(
                    tuple(sheet.cell(1, column).value for column in range(1, 11)),
                    REVIEW_COLUMNS,
                )
                self.assertEqual(sheet.max_row, 2)
                self.assertEqual(sheet.cell(2, 2).value, 4)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()

